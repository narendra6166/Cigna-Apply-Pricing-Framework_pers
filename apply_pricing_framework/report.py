"""Write business-readable reports without request PII."""
from __future__ import annotations

import csv
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from .engine import Outcome

HEADERS = [
    "Scenario ID",
    "Scenario Name",
    "Category",
    "Payload Name",
    "Expected Result",
    "Actual Result",
    "Comments",
    "Status",
    "Execution Date & Time",
    "Design References",
]


def _row(outcome: Outcome) -> list[str]:
    return [
        outcome.scenario_id,
        outcome.scenario_name,
        outcome.category,
        outcome.payload_name,
        outcome.expected,
        outcome.actual,
        outcome.comments,
        outcome.status,
        outcome.executed_at,
        outcome.design_refs,
    ]


def write_csv(path: Path, outcomes: list[Outcome]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(_row(outcome) for outcome in outcomes)


def write_markdown(
    path: Path, outcomes: list[Outcome], unmapped: list[dict[str, Any]], report_time: str
) -> None:
    counts = Counter(outcome.status for outcome in outcomes)
    lines = [
        "# Apply Pricing execution report",
        "",
        f"Report generated: **{report_time}**",
        f"Active design scenarios: **{len({outcome.scenario_id for outcome in outcomes})}** · "
        f"Payload executions: **{len(outcomes)}** · "
        f"Pass: **{counts['Pass']}** · Fail: **{counts['Fail']}** · "
        f"Partial Pass: **{counts['Partial Pass']}** · Blocked: **{counts['Blocked']}** · "
        f"Not Run: **{counts['Not Run']}**",
        "",
        "Expected results are evaluated from `apply_pricing_framework/rules/expected-results.json`. "
        "Unmapped collection requests are listed separately and excluded from scenario pass/fail totals. "
        "Blocked means the payload did not meet its scenario preconditions, execution data was unavailable, "
        "or its design condition is unresolved; it is not counted as a business failure.",
        "",
        "| Scenario ID | Scenario | Payload | Status | Expected | Actual | Comments |",
        "|---|---|---|---|---|---|---|",
    ]
    for outcome in outcomes:
        values = [
            outcome.scenario_id,
            outcome.scenario_name,
            outcome.payload_name,
            outcome.status,
            outcome.expected,
            outcome.actual,
            outcome.comments,
        ]
        lines.append("| " + " | ".join(_md_cell(value) for value in values) + " |")
    lines.extend(["", "## Bruno requests outside the active design", ""])
    if unmapped:
        lines.extend(
            [
                "| Payload | Bruno path | Response | Handling |",
                "|---|---|---|---|",
            ]
        )
        for item in unmapped:
            lines.append(
                "| "
                + " | ".join(
                    _md_cell(str(item[key]))
                    for key in ("payload_name", "path", "response_status", "note")
                )
                + " |"
            )
    else:
        lines.append("None.")
    lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def _md_cell(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\r", " ").replace("\n", "<br>")


def write_excel(
    path: Path,
    outcomes: list[Outcome],
    unmapped: list[dict[str, Any]],
    catalog: dict[str, Any],
    report_time: str,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError(
            "Excel output requires openpyxl. Install it with "
            "'python -m pip install -r apply_pricing_framework\\requirements.txt'."
        ) from exc

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["CCAC-4226 / PTTC-4106 — Apply Pricing"])
    summary.merge_cells("A1:B1")
    summary["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="17365D")
    counts = Counter(outcome.status for outcome in outcomes)
    for key, value in [
        ("Report generated", report_time),
        ("Scenario source", catalog["sources"]["scenario_list"]),
        ("Design source", catalog["sources"]["pricing_design"]),
        ("Scenario outcomes", len(outcomes)),
        ("Pass", counts["Pass"]),
        ("Fail", counts["Fail"]),
        ("Partial Pass", counts["Partial Pass"]),
        ("Blocked", counts["Blocked"]),
        ("Not Run", counts["Not Run"]),
        ("Unmapped Bruno requests (not scored)", len(unmapped)),
    ]:
        summary.append([key, value])
    summary.append([])
    summary.append(["How to read this report", "Blocked inputs / unavailable data are not Pass or Fail; see Actual Result for the reason."])
    summary.append(["Design conflicts", "See the Design Decisions sheet and apply_pricing_framework/design/TRACEABILITY.md."])
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 110
    for cell in summary[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    results = workbook.create_sheet("Scenario Results")
    results.append(HEADERS)
    for outcome in outcomes:
        results.append(_row(outcome))
    widths = [18, 43, 25, 48, 80, 90, 85, 16, 32, 72]
    for index, width in enumerate(widths, 1):
        results.column_dimensions[chr(64 + index)].width = width
    for cell in results[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    colors = {
        "Pass": "E2F0D9",
        "Fail": "F4CCCC",
        "Partial Pass": "FCE5CD",
        "Blocked": "D9E1F2",
        "Not Run": "E7E6E6",
    }
    for row in results.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[7].fill = PatternFill("solid", fgColor=colors.get(row[7].value, "FFFFFF"))
    if results.max_row > 1:
        table = Table(displayName="ApplyPricingScenarioResults", ref=f"A1:J{results.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        results.add_table(table)
    results.freeze_panes = "A2"

    unmapped_sheet = workbook.create_sheet("Unmapped Payloads")
    unmapped_sheet.append(["Payload Name", "Bruno Path", "HTTP Response", "Handling"])
    for item in unmapped:
        unmapped_sheet.append(
            [item["payload_name"], item["path"], item["response_status"], item["note"]]
        )
    for index, width in enumerate([48, 72, 18, 110], 1):
        unmapped_sheet.column_dimensions[chr(64 + index)].width = width
    for cell in unmapped_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for row in unmapped_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    decisions = workbook.create_sheet("Design Decisions")
    decisions.append(["ID", "Topic", "Design comparison", "Treatment"])
    for item in catalog.get("design_conflicts", []):
        decisions.append(
            [item["id"], item["topic"], item.get("scenario_source", "") + "\n" + item.get("detailed_design", item.get("pricing_design", item.get("bruno_payloads", ""))), item["treatment"]]
        )
    for index, width in enumerate([16, 32, 100, 100], 1):
        decisions.column_dimensions[chr(64 + index)].width = width
    for cell in decisions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for row in decisions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    except OSError as exc:
        raise PermissionError(
            f"Could not update {path}; close the workbook if it is open and retry."
        ) from exc
    finally:
        if temporary.exists():
            temporary.unlink()


def write_word(
    path: Path,
    outcomes: list[Outcome],
    unmapped: list[dict[str, Any]],
    catalog: dict[str, Any],
    report_time: str,
) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError(
            "Word output requires python-docx. Install it with "
            "'python -m pip install -r apply_pricing_framework\\requirements.txt'."
        ) from exc

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    title = document.add_heading("Apply Pricing Execution Report", level=0)
    title.runs[0].font.color.rgb = RGBColor(23, 54, 93)
    document.add_paragraph(
        f"Scenario source: {catalog['sources']['scenario_list']} | "
        f"Design source: {catalog['sources']['pricing_design']}"
    )

    counts = Counter(outcome.status for outcome in outcomes)
    summary = document.add_table(rows=0, cols=2)
    summary.style = "Light Shading Accent 1"
    for label, value in [
        ("Report generated", report_time),
        ("Design scenarios", len({outcome.scenario_id for outcome in outcomes})),
        ("Payload executions", len(outcomes)),
        ("Pass", counts["Pass"]),
        ("Fail", counts["Fail"]),
        ("Partial Pass", counts["Partial Pass"]),
        ("Blocked", counts["Blocked"]),
        ("Not Run", counts["Not Run"]),
        ("Unmapped Bruno requests (not scored)", len(unmapped)),
    ]:
        cells = summary.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        cells[0].paragraphs[0].runs[0].bold = True

    document.add_paragraph(
        "Expected results are evaluated from the centralized scenario catalog. "
        "Blocked means the payload did not meet its preconditions, execution data was unavailable, "
        "or a design condition is unresolved; it is not counted as a business failure."
    )
    document.add_heading("Scenario Results", level=1)
    status_colors = {
        "Pass": "548235",
        "Fail": "C00000",
        "Partial Pass": "C65911",
        "Blocked": "4472C4",
        "Not Run": "666666",
    }
    for outcome in outcomes:
        heading = document.add_heading(
            f"{outcome.scenario_id} — {outcome.scenario_name}", level=2
        )
        for run in heading.runs:
            run.font.color.rgb = RGBColor.from_string(status_colors.get(outcome.status, "17365D"))
        details = document.add_table(rows=0, cols=2)
        details.style = "Light Shading Accent 1"
        for label, value in [
            ("Category", outcome.category),
            ("Payload", outcome.payload_name),
            ("Status", outcome.status),
            ("Execution Date & Time", outcome.executed_at),
            ("Expected Result", outcome.expected),
            ("Actual Result", outcome.actual),
            ("Comments", outcome.comments),
            ("Design References", outcome.design_refs),
        ]:
            cells = details.add_row().cells
            cells[0].text = label
            cells[1].text = value or ""
            cells[0].paragraphs[0].runs[0].bold = True
        document.add_paragraph()

    document.add_heading("Unmapped Bruno Requests", level=1)
    if unmapped:
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, label in zip(
            table.rows[0].cells, ["Payload Name", "Bruno Path", "HTTP Response", "Handling"]
        ):
            cell.text = label
            cell.paragraphs[0].runs[0].bold = True
        for item in unmapped:
            cells = table.add_row().cells
            for cell, key in zip(
                cells, ("payload_name", "path", "response_status", "note")
            ):
                cell.text = str(item.get(key, ""))
    else:
        document.add_paragraph("None.")

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        document.save(temporary)
        temporary.replace(path)
    except PermissionError as exc:
        raise PermissionError(
            f"Could not update {path}; close the Word report if it is open and retry."
        ) from exc
    finally:
        if temporary.exists():
            temporary.unlink()
