from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZipFile

import apply_pricing_framework.__main__ as framework_cli
from bruno.update_payload_from_mainframe import parse_export, update_payload
from apply_pricing_framework.engine import CatalogError, evaluate, load_catalog, run_catalog
from apply_pricing_framework.report import write_csv, write_excel, write_markdown, write_word

ROOT = Path(__file__).resolve().parents[2]
CATALOG_PATH = ROOT / "apply_pricing_framework" / "rules" / "expected-results.json"


class MainframePayloadSyncTests(unittest.TestCase):
    def test_parses_source_with_missing_array_terminator(self) -> None:
        from bruno.update_payload_from_mainframe import parse_export

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "source.yaml"
            source_path.write_text(
                '{"query": [{"value": 1}\n# next query\n{"next": [{"ok": true}]}',
                encoding="utf-8",
            )
            self.assertEqual(
                parse_export(source_path),
                [{"query": [{"value": 1}]}, {"next": [{"ok": True}]}],
            )

    def test_updates_matching_request_fields_from_concatenated_export_blocks(self) -> None:
        source = """# Header extract
{"header query": [{
  "P20_GCDOC_NO": 12345,
  "REQUESTDATE": 20260618,
  "EMPLOYEESSN": 111222333,
  "PATIENTSUFFIX": 3,
  "PATIENTDOB": 20010102,
  "SOC": 1234,
  "CLIENTID": 5678,
  "ACCOUNTNUMBER": 9012,
  "BENOPT": "OAP",
  "SCHEDULEATTR": "APCOP",
  "PRICINGTYPE": "APC",
  "DRGVERSION": 26,
  "USEDDRG": "    ",
  "ALLOW_CHG": 25.5,
  "GRP_CHG": 30.0,
  "USAGE": "A ",
  "RATETYPE": "LN-DENIED-015,P10",
  "BASED_ON": "SI-N,AP",
  "APC_NUM": "5023"
}]}
# Service extract
{"service query": [{
  "SERVICEFROMDATE": 20260618,
  "SERVICETHRUDATE": 20260618,
  "PROCEDURECODE": "87086",
  "REVENUECODE": -403,
  "CHARGEAMT": 30.0,
  "MODIFIER1": null
}]}
"""
        request = {
            "serviceReferenceId": "AP-001_111_2",
            "requestDate": "20260101",
            "memberInformation": {
                "employeeSsn": "",
                "patientSuffix": "",
            },
            "planInformation": {"soc": ""},
            "providers": {},
            "claimInformation": {
                "claimLevel": {
                    "claimCcrKey": "111_2",
                    "claimReceiveDate": "20260101",
                    "pricingDetails": {
                        "allowedCharges": 0,
                        "pricingRules": [{
                            "pdBasedonCd": [{"pdBasedonCode": ""}, {"pdBasedonCode": ""}],
                            "pdRateTypeCd": [{"pdRateTypeCode": ""}],
                            "pdUsage": "",
                            "pdLineNum": 1,
                            "pdOriginalLineNum": 1,
                        }],
                    },
                },
                "serviceLevel": [{
                    "serviceFromDate": "20260101",
                    "serviceThruDate": "20260101",
                    "procedureCode": "00000",
                    "revenueCode": "00000",
                    "chargeAmt": 0,
                    "modifier1": "OLD",
                    "pricingArea": {
                        "pdAllowed": 0,
                        "pdGroupCharge": 0,
                        "pdMaxUnits": 0,
                        "pdApcNum": 0,
                    },
                    "networkArea": {},
                }],
            },
        }
        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "source.yaml"
            payload_path = Path(directory) / "request.bru"
            source_path.write_text(source, encoding="utf-8")
            original_root_fields = set(request)
            original_line_fields = set(request["claimInformation"]["serviceLevel"][0])
            payload_path.write_text(
                "meta {\n  name: AP-001 example\n}\nbody:json {\n  "
                + json.dumps(request)
                + "\n}\n",
                encoding="utf-8",
            )

            blocks = parse_export(source_path)
            lines_updated, source_lines = update_payload(payload_path, source_path)
            body_text = payload_path.read_text(encoding="utf-8")
            body_start = body_text.index("body:json {") + len("body:json {")
            body_start = body_text.index("{", body_start)
            body_end = body_text.rindex("}")
            updated = json.loads(body_text[body_start:body_end].strip())

        self.assertEqual(len(blocks), 2)
        self.assertEqual((lines_updated, source_lines), (1, 1))
        self.assertEqual(set(updated), original_root_fields)
        self.assertEqual(updated["memberInformation"]["employeeSsn"], "111222333")
        self.assertEqual(updated["memberInformation"]["patientSuffix"], "03")
        self.assertEqual(updated["claimInformation"]["claimLevel"]["claimCcrKey"], "12345_2")
        line = updated["claimInformation"]["serviceLevel"][0]
        self.assertEqual(set(line), original_line_fields)
        self.assertEqual(line["procedureCode"], "87086")
        self.assertEqual(line["revenueCode"], "00403")
        self.assertEqual(line["modifier1"], "")
        self.assertEqual(line["pricingArea"]["pdAllowed"], 25.5)
        self.assertEqual(
            updated["claimInformation"]["claimLevel"]["pricingDetails"]["pricingRules"][0]
            ["pdBasedonCd"],
            [{"pdBasedonCode": "SI-N"}, {"pdBasedonCode": "AP"}],
        )

    def test_rebuilds_service_lines_and_rules_to_match_source_count(self) -> None:
        source_rows = [{
            "P20_GCDOC_NO": 12345,
            "EMPLOYEESSN": 111222333,
        }, {
            "PROCEDURECODE": "87086",
        }, {
            "PROCEDURECODE": "80050",
            "CHARGEAMT": 20,
        }]
        from bruno.update_payload_from_mainframe import update_request

        request = {
            "memberInformation": {},
            "planInformation": {},
            "providers": {},
            "claimInformation": {
                "claimLevel": {
                    "pricingDetails": {
                        "pricingRules": [{
                            "pdBasedonCd": [{"pdBasedonCode": ""}],
                            "pdRateTypeCd": [{"pdRateTypeCode": ""}],
                            "pdUsage": "",
                            "pdLineNum": 1,
                            "pdOriginalLineNum": 1,
                        }],
                    },
                },
                "serviceLevel": [{
                    "procedureCode": "old",
                    "lineNum": "0001",
                    "pricingArea": {"pdLineNum": 1},
                    "networkArea": {},
                }],
            },
        }
        update_request(request, [{"query": source_rows}])
        self.assertEqual(
            [line["procedureCode"] for line in request["claimInformation"]["serviceLevel"]],
            ["87086", "80050"],
        )
        rules = request["claimInformation"]["claimLevel"]["pricingDetails"]["pricingRules"]
        self.assertEqual(len(rules), 2)
        self.assertEqual([rule["pdLineNum"] for rule in rules], [1, 2])
        self.assertEqual(set(rules[0]), set(rules[1]))
        self.assertEqual(
            set(request["claimInformation"]["serviceLevel"][1]),
            set(request["claimInformation"]["serviceLevel"][0]),
        )


class CatalogTests(unittest.TestCase):
    def test_catalog_matches_all_current_design_ids(self) -> None:
        catalog = load_catalog(CATALOG_PATH)
        self.assertEqual(len(catalog["scenarios"]), 36)
        self.assertEqual(len({item["id"] for item in catalog["scenarios"]}), len(catalog["scenarios"]))
        self.assertTrue(all(item["design_refs"] for item in catalog["scenarios"]))
        doc_path = CATALOG_PATH.parent.parent / "design" / (
            "CCAC-4226-PTTC-4106-Apply-Pricing-System-Test-Scenarios-10-16-2026.docx"
        )
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with ZipFile(doc_path) as archive:
            document = ElementTree.fromstring(archive.read("word/document.xml"))
        design_ids = {
            " ".join(text.text or "" for text in row.findall("./w:tc[1]//w:t", namespace)).strip()
            for row in document.findall(".//w:tbl/w:tr", namespace)
        } - {"Scenario ID"}
        self.assertEqual({item["id"] for item in catalog["scenarios"]}, design_ids)
        self.assertEqual(len(design_ids), 36)

    def test_duplicate_payload_mapping_is_rejected(self) -> None:
        catalog = {
            "schema_version": 1,
            "scenarios": [
                {
                    "id": scenario_id,
                    "name": scenario_id,
                    "category": "Test",
                    "expected": "Expected",
                    "payloads": ["same.bru"],
                    "conditions": [],
                    "checks": [{"op": "response_message_code_in"}],
                    "design_refs": ["unit test"],
                }
                for scenario_id in ("A", "B")
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "catalog.json"
            path.write_text(json.dumps(catalog), encoding="utf-8")
            with self.assertRaisesRegex(CatalogError, "mapped to more than one scenario"):
                load_catalog(path)


class EvaluationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.catalog = {
            "sources": {"scenario_list": "scenarios.docx", "pricing_design": "pricing.pdf"},
            "execution_blockers": {
                "response_message_codes": ["1301"],
                "response_message_text_contains": [],
                "http_statuses": [500, 502, 503, 504],
            },
            "scenarios": [
                {
                    "id": "SAMPLE",
                    "name": "Successful sample",
                    "category": "Unit tests",
                    "payloads": ["01-sample/sample.bru"],
                    "expected": "Claim and line succeed.",
                    "conditions": [],
                    "checks": [
                        {"op": "response_message_code_in", "values": ["200"], "severity": "critical"}
                    ],
                    "design_refs": ["Unit test fixture"],
                }
            ],
        }
        self.result = {
            "test": {"filename": "01-sample/sample.bru"},
            "request": {"data": {"claimInformation": {"claimLevel": {"pricingDetails": {}}, "serviceLevel": []}}},
            "response": {
                "status": 200,
                "headers": {"Date": "Wed, 16 Oct 2026 12:00:00 GMT"},
                "data": {"message": {"code": "200"}, "claimLevel": {"claimLevelGroupMethod": "AP"}, "lineItems": []},
            },
        }

    def test_result_passes_and_execution_date_comes_from_response(self) -> None:
        outcomes, unmapped = run_catalog(self.catalog, [self.result], "report time")
        self.assertEqual(outcomes[0].status, "Pass")
        self.assertEqual(outcomes[0].executed_at, "Wed, 16 Oct 2026 12:00:00 GMT")
        self.assertEqual(unmapped, [])

    def test_missing_result_is_not_run(self) -> None:
        outcomes, _ = run_catalog(self.catalog, [], "report time")
        self.assertEqual(outcomes[0].status, "Not Run")

    def test_unmapped_request_is_separate(self) -> None:
        result = dict(self.result)
        result["test"] = {"filename": "99-other/other.bru"}
        outcomes, unmapped = run_catalog(self.catalog, [result], "report time")
        self.assertEqual(outcomes[0].status, "Not Run")
        self.assertEqual(len(unmapped), 1)

    def test_failed_check_has_mismatch_comment(self) -> None:
        response = self.result["response"]["data"]
        response["message"]["code"] = "7011"
        outcomes, _ = run_catalog(self.catalog, [self.result], "report time")
        self.assertEqual(outcomes[0].status, "Fail")
        self.assertIn("expected one of ['200']; got 7011", outcomes[0].comments)

    def test_passed_check_has_no_comment(self) -> None:
        outcomes, _ = run_catalog(self.catalog, [self.result], "report time")
        self.assertEqual(outcomes[0].status, "Pass")
        self.assertEqual(outcomes[0].comments, "")

    def test_reports_contain_outcome_without_request_payload(self) -> None:
        outcomes, unmapped = run_catalog(self.catalog, [self.result], "report time")
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "result.csv"
            md_path = Path(directory) / "result.md"
            xlsx_path = Path(directory) / "result.xlsx"
            docx_path = Path(directory) / "result.docx"
            write_csv(csv_path, outcomes)
            write_markdown(md_path, outcomes, unmapped, "report time")
            write_excel(xlsx_path, outcomes, unmapped, self.catalog, "report time")
            write_word(docx_path, outcomes, unmapped, self.catalog, "report time")
            self.assertIn("Scenario ID", csv_path.read_text(encoding="utf-8-sig"))
            self.assertIn("Successful sample", md_path.read_text(encoding="utf-8"))
            self.assertNotIn("claimInformation", md_path.read_text(encoding="utf-8"))
            from docx import Document

            word = Document(docx_path)
            word_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
            word_text += "\n" + "\n".join(
                cell.text for table in word.tables for row in table.rows for cell in row.cells
            )
            self.assertIn("Successful sample", word_text)
            self.assertIn("Claim and line succeed.", word_text)
            self.assertIn("Scenario Results", word_text)
            self.assertNotIn("claimInformation", word_text)
            from openpyxl import load_workbook

            workbook = load_workbook(xlsx_path, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Summary", "Scenario Results", "Unmapped Payloads", "Design Decisions"],
            )
            self.assertEqual(workbook["Scenario Results"].max_row, 2)
            workbook.close()

    def test_add_scenario_creates_blocked_scaffold(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            project_root = Path(directory)
            payload = project_root / "bruno" / "01-sample" / "new-sample.bru"
            payload.parent.mkdir(parents=True)
            payload.write_text("request", encoding="utf-8")
            catalog_path = project_root / "catalog.json"
            catalog_path.write_text(
                json.dumps(load_catalog(CATALOG_PATH)),
                encoding="utf-8",
            )
            with patch.object(framework_cli, "PROJECT_ROOT", project_root):
                exit_code = framework_cli.main(
                    [
                        "add-scenario",
                        "--id",
                        "AP-NEW-TEST",
                        "--name",
                        "New test scenario",
                        "--payload",
                        "01-sample/new-sample.bru",
                        "--catalog",
                        str(catalog_path),
                    ]
                )
            self.assertEqual(exit_code, 0)
            updated = load_catalog(catalog_path)
            scenario = next(item for item in updated["scenarios"] if item["id"] == "AP-NEW-TEST")
            result = {
                "test": {"filename": "01-sample/new-sample.bru"},
                "request": {"data": {"claimInformation": {"claimLevel": {}, "serviceLevel": []}}},
                "response": {
                    "status": 200,
                    "data": {"message": {"code": "200"}, "claimLevel": {}, "lineItems": []},
                },
            }
            from apply_pricing_framework.engine import evaluate

            outcome = evaluate(scenario, result, updated, "report time")
            self.assertEqual(outcome.status, "Blocked")

    def test_plan_limit_case_rate_expected_amounts_reconcile(self) -> None:
        scenario = {
            "id": "AP-031",
            "name": "Finite plan limit",
            "category": "Case Rate",
            "payloads": ["case/AP-031.bru"],
            "expected": "Use finite limit for included line basis.",
            "conditions": [
                {"op": "group_method_equals", "value": "10"},
                {"op": "plan_limit_relation_present", "value": "lt_billed"},
            ],
            "checks": [
                {"op": "response_group_method_equals", "value": "AR"},
                {"op": "case_rate_amounts_reconcile"},
            ],
            "design_refs": ["PDF p.4"],
        }
        request = {
            "claimInformation": {
                "claimLevel": {"pricingDetails": {"groupMethod": "10", "allowedCharges": 9600}},
                "serviceLevel": [
                    {"lineNum": "0001", "chargeAmt": 9000, "planCoveredExpenseLimit": 2000,
                     "pricingArea": {"pdAllowed": 7200, "clExcluded": "FALSE"}},
                    {"lineNum": "0002", "chargeAmt": 3000,
                     "pricingArea": {"pdAllowed": 2400, "clExcluded": "FALSE"}},
                ],
            }
        }
        response = {
            "message": {"code": "200"},
            "claimLevel": {"claimLevelGroupMethod": "AR", "claimLevelTotalNotCoveredAmount": 7000},
            "lineItems": [
                {"lineNumber": "0001", "message": {"code": "200"}, "totalNotCoveredAmount": 7000},
                {"lineNumber": "0002", "message": {"code": "200"}, "totalNotCoveredAmount": 0},
            ],
        }
        result = {
            "test": {"filename": "case/AP-031.bru"},
            "request": {"data": request},
            "response": {"status": 200, "data": response},
        }
        outcome = evaluate(scenario, result, {"execution_blockers": {}}, "report time")
        self.assertEqual(outcome.status, "Pass")
        response["claimLevel"]["claimLevelTotalNotCoveredAmount"] = 2400
        outcome = evaluate(scenario, result, {"execution_blockers": {}}, "report time")
        self.assertEqual(outcome.status, "Fail")
        self.assertIn("claim NC expected 7000.00", outcome.comments)

    def test_unlimited_plan_limit_uses_billed_basis(self) -> None:
        scenario = {
            "id": "AP-033",
            "name": "Unlimited plan limit",
            "category": "Case Rate",
            "payloads": ["case/AP-033.bru"],
            "expected": "All-9s means unlimited.",
            "conditions": [{"op": "plan_limit_unlimited_present", "value": True}],
            "checks": [{"op": "case_rate_amounts_reconcile"}],
            "design_refs": ["PDF p.4"],
        }
        request = {
            "claimInformation": {
                "claimLevel": {"pricingDetails": {"allowedCharges": 9600}},
                "serviceLevel": [
                    {"lineNum": "0001", "chargeAmt": 9000, "planCoveredExpenseLimit": 999999999.99,
                     "pricingArea": {"pdAllowed": 7200, "clExcluded": "FALSE"}},
                    {"lineNum": "0002", "chargeAmt": 3000,
                     "pricingArea": {"pdAllowed": 2400, "clExcluded": "FALSE"}},
                ],
            }
        }
        response = {
            "message": {"code": "200"},
            "claimLevel": {"claimLevelTotalNotCoveredAmount": 2400},
            "lineItems": [
                {"lineNumber": "0001", "message": {"code": "200"}, "totalNotCoveredAmount": 1800},
                {"lineNumber": "0002", "message": {"code": "200"}, "totalNotCoveredAmount": 600},
            ],
        }
        result = {
            "test": {"filename": "case/AP-033.bru"},
            "request": {"data": request},
            "response": {"status": 200, "data": response},
        }
        outcome = evaluate(scenario, result, {"execution_blockers": {}}, "report time")
        self.assertEqual(outcome.status, "Pass")


if __name__ == "__main__":
    unittest.main()
